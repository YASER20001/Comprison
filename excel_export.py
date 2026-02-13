"""Color-coded Excel export module using openpyxl."""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Color configuration
COLORS = {
    "added":     {"bg": "C6EFCE", "font": "006100", "label": "ADDED"},
    "deleted":   {"bg": "FFC7CE", "font": "9C0006", "label": "DELETED"},
    "modified":  {"bg": "FFE0B2", "font": "E65100", "label": "MODIFIED"},
    "unchanged": {"bg": "E0E0E0", "font": "424242", "label": "NO CHANGE"},
}
HEADER_BG = "003366"
HEADER_FONT = "FFFFFF"


def _stringify(value):
    if value is None:
        return ""
    s = str(value)
    return "" if s == "nan" else s


def export_comparison(result, before_name, after_name):
    """
    Export comparison result to a color-coded Excel workbook.

    Returns an io.BytesIO buffer containing the .xlsx file.
    """
    wb = Workbook()

    # ── Sheet 1: Summary ──
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.sheet_properties.tabColor = "003366"

    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 35

    title_font = Font(bold=True, size=16, color="003366")
    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"].value = "FILE COMPARISON REPORT"
    ws_summary["A1"].font = title_font

    summary = result["summary"]
    info_rows = [
        ("Before File:", before_name),
        ("After File:", after_name),
        ("Compared At:", result["comparedAt"]),
        ("Key Column:", result["keyColumn"]),
    ]
    for i, (label, value) in enumerate(info_rows, start=3):
        ws_summary.cell(row=i, column=1, value=label).font = Font(bold=True, size=11)
        ws_summary.cell(row=i, column=2, value=value)

    # Stats header
    r = 8
    ws_summary.merge_cells(f"A{r}:B{r}")
    ws_summary.cell(row=r, column=1, value="CHANGE SUMMARY").font = Font(bold=True, size=13, color="003366")

    stat_rows = [
        ("Total Records",       summary["total"],     None),
        ("Added (New)",         summary["added"],     COLORS["added"]["bg"]),
        ("Deleted (Removed)",   summary["deleted"],   COLORS["deleted"]["bg"]),
        ("Modified (Changed)",  summary["modified"],  COLORS["modified"]["bg"]),
        ("Unchanged",           summary["unchanged"], COLORS["unchanged"]["bg"]),
    ]
    for i, (label, value, bg) in enumerate(stat_rows, start=10):
        c1 = ws_summary.cell(row=i, column=1, value=label)
        c2 = ws_summary.cell(row=i, column=2, value=value)
        c1.font = Font(bold=True, size=11)
        c2.font = Font(bold=True, size=11)
        if bg:
            fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            c1.fill = fill
            c2.fill = fill

    # ── Sheet 2: Comparison Details ──
    ws_detail = wb.create_sheet("Comparison Details")
    ws_detail.sheet_properties.tabColor = "0066CC"

    columns = result["columns"]
    headers = ["Status", "Key"] + columns
    header_fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    header_font = Font(bold=True, color=HEADER_FONT, size=11)
    thin_border = Border(bottom=Side(style="thin", color="000000"))

    for col_idx, header in enumerate(headers, start=1):
        cell = ws_detail.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws_detail.column_dimensions[get_column_letter(col_idx)].width = max(15, len(header) + 5)

    hair_border = Border(bottom=Side(style="hair", color="CCCCCC"))

    for row_idx, comp_row in enumerate(result["rows"], start=2):
        status = comp_row["status"]
        color = COLORS[status]
        row_fill = PatternFill(start_color=color["bg"], end_color=color["bg"], fill_type="solid")
        row_font = Font(color=color["font"])

        data = comp_row["afterData"] or comp_row["beforeData"] or {}

        # Status cell
        cell = ws_detail.cell(row=row_idx, column=1, value=color["label"])
        cell.fill = row_fill
        cell.font = Font(bold=True, color=color["font"])
        cell.border = hair_border

        # Key cell
        cell = ws_detail.cell(row=row_idx, column=2, value=comp_row["keyValue"])
        cell.fill = row_fill
        cell.font = row_font
        cell.border = hair_border

        # Data cells
        changed_cols = {c["column"] for c in comp_row["changes"]}
        for col_idx, col_name in enumerate(columns, start=3):
            cell = ws_detail.cell(row=row_idx, column=col_idx, value=_stringify(data.get(col_name, "")))
            cell.fill = row_fill
            cell.font = row_font
            cell.border = hair_border

            # Highlight changed cells in modified rows
            if status == "modified" and col_name in changed_cols:
                change = next(c for c in comp_row["changes"] if c["column"] == col_name)
                cell.value = f"{_stringify(change['oldValue'])} → {_stringify(change['newValue'])}"
                cell.font = Font(bold=True, color="E65100")
                cell.fill = PatternFill(start_color="FFD180", end_color="FFD180", fill_type="solid")

    # Auto-filter and freeze
    ws_detail.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(result['rows']) + 1}"
    ws_detail.freeze_panes = "C2"

    # ── Sheet 3: Changes Only ──
    ws_changes = wb.create_sheet("Changes Only")
    ws_changes.sheet_properties.tabColor = "F58220"

    change_headers = ["Status", "Key", "Column", "Before Value", "After Value"]
    for col_idx, header in enumerate(change_headers, start=1):
        cell = ws_changes.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_changes.column_dimensions["A"].width = 15
    ws_changes.column_dimensions["B"].width = 25
    ws_changes.column_dimensions["C"].width = 25
    ws_changes.column_dimensions["D"].width = 30
    ws_changes.column_dimensions["E"].width = 30

    change_row = 2
    for comp_row in result["rows"]:
        status = comp_row["status"]
        if status == "unchanged":
            continue

        color = COLORS[status]
        row_fill = PatternFill(start_color=color["bg"], end_color=color["bg"], fill_type="solid")
        row_font = Font(color=color["font"])

        if status == "modified":
            for change in comp_row["changes"]:
                for ci, val in enumerate([
                    color["label"],
                    comp_row["keyValue"],
                    change["column"],
                    _stringify(change["oldValue"]),
                    _stringify(change["newValue"]),
                ], start=1):
                    cell = ws_changes.cell(row=change_row, column=ci, value=val)
                    cell.fill = row_fill
                    cell.font = row_font
                change_row += 1
        else:
            values = [
                color["label"],
                comp_row["keyValue"],
                "—",
                "(entire row)" if status == "deleted" else "",
                "(entire row)" if status == "added" else "",
            ]
            for ci, val in enumerate(values, start=1):
                cell = ws_changes.cell(row=change_row, column=ci, value=val)
                cell.fill = row_fill
                cell.font = row_font
            change_row += 1

    ws_changes.freeze_panes = "A2"

    # Write to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
